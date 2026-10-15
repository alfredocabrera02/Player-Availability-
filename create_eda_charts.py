"""
create_eda_charts.py
Generate 4 EDA charts for the presentation EDA slide.
Run: python create_eda_charts.py
"""
import warnings
warnings.filterwarnings("ignore")
import os
import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches

BASE = r"C:\Users\fredo\500 Project"

NAVY  = "#1B2E5E"
RED   = "#B9182F"
GREEN = "#167A3C"
BLUE  = "#1D6FD8"
AMBER = "#D97706"
LGRAY = "#E2E8F0"
DGRAY = "#475569"
WHITE = "#FFFFFF"
OFF_W = "#F8FAFC"


def safe_num(v):
    if v is None:
        return np.nan
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s.startswith("="):
        try:
            return float(eval(s[1:]))
        except Exception:
            return np.nan
    try:
        return float(s)
    except (TypeError, ValueError):
        return np.nan


print("Loading Row per Player...")
wb = openpyxl.load_workbook(os.path.join(BASE, "data columns v2.xlsx"), data_only=True)
ws = wb["Row per Player"]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
rows = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)]
df = pd.DataFrame(rows, columns=headers)

df["match_date"] = pd.to_datetime(df["match_date"])
df["season"] = df["match_date"].dt.year

num_cols = [
    "minutes_on_date", "minutes_last_match", "minutes_last_7_days",
    "minutes_last_14_days", "minutes_last_30_days", "matches_last_14_days",
    "consecutive_starts", "played_last_match", "days_since_last_match",
    "missed_matches_last_30_days", "days_since_last_absence",
    "injuries_last_30_days", "injuries_last_60_days",
    "player_unavailable_next_match",
]
for col in num_cols:
    if col in df.columns:
        df[col] = df[col].apply(safe_num)

labeled = df[df["player_unavailable_next_match"].notna()].copy()
labeled["player_unavailable_next_match"] = labeled["player_unavailable_next_match"].astype(int)
labeled = labeled.dropna(subset=[
    "minutes_last_match", "minutes_last_7_days",
    "minutes_last_14_days", "minutes_last_30_days",
    "days_since_last_match", "played_last_match",
])
labeled["days_since_last_absence"] = labeled["days_since_last_absence"].fillna(999)
print(f"Labeled rows: {len(labeled)}  |  Positive rate: {labeled['player_unavailable_next_match'].mean()*100:.1f}%")

plt.rcParams.update({
    "font.family": "sans-serif",
    "axes.spines.top": False,
    "axes.spines.right": False,
    "figure.facecolor": WHITE,
    "axes.facecolor": OFF_W,
    "grid.color": LGRAY,
    "grid.alpha": 0.7,
})


# ── Chart 1: Absence rate by position ─────────────────────────────────────────
print("Chart 1: Absence rate by position...")
pos_g = (labeled.groupby("position", observed=True)
         .agg(rate=("player_unavailable_next_match", "mean"),
              n=("player_unavailable_next_match", "count"))
         .reset_index()
         .sort_values("rate", ascending=True))

fig1, ax1 = plt.subplots(figsize=(5.5, 2.8))
colors1 = [RED if r > 0.45 else AMBER if r > 0.30 else GREEN for r in pos_g["rate"]]
bars1 = ax1.barh(pos_g["position"], pos_g["rate"] * 100,
                  color=colors1, edgecolor=WHITE, height=0.62, zorder=3)
ax1.axvline(37, color=NAVY, ls="--", lw=1.5, label="Squad avg (37%)", zorder=4)
for bar, r, n in zip(bars1, pos_g["rate"], pos_g["n"]):
    ax1.text(bar.get_width() + 0.8, bar.get_y() + bar.get_height() / 2,
             f"{r*100:.0f}%  n={n}", va="center", fontsize=8, color=DGRAY)
ax1.set_xlabel("Absence Rate (%)", fontsize=9, color=DGRAY)
ax1.set_xlim(0, 90)
ax1.set_title("Absence Rate by Position", fontsize=11, color=NAVY, fontweight="bold", pad=6)
ax1.legend(fontsize=8, framealpha=0.9, loc="lower right")
ax1.grid(axis="x", zorder=0)
ax1.tick_params(labelsize=9)
plt.tight_layout(pad=0.6)
p1 = os.path.join(BASE, "eda_absence_by_position.png")
fig1.savefig(p1, dpi=150, bbox_inches="tight", facecolor=WHITE)
plt.close()
print(f"  Saved: {p1}")


# ── Chart 2: Key features — available vs unavailable ──────────────────────────
print("Chart 2: Workload distributions...")
avail   = labeled[labeled["player_unavailable_next_match"] == 0]
unavail = labeled[labeled["player_unavailable_next_match"] == 1]

fig2, axes2 = plt.subplots(1, 2, figsize=(5.5, 2.8))
fig2.patch.set_facecolor(WHITE)
for ax, feat, title in zip(
        axes2,
        ["consecutive_starts", "minutes_last_7_days"],
        ["Consecutive Starts", "Minutes — Last 7 Days"]):
    d0 = avail[feat].dropna().values
    d1 = unavail[feat].dropna().values
    parts = ax.violinplot([d0, d1], positions=[0, 1],
                           showmedians=True, showextrema=False, widths=0.7)
    for pc, c in zip(parts["bodies"], [GREEN, RED]):
        pc.set_facecolor(c)
        pc.set_alpha(0.72)
    parts["cmedians"].set_color(WHITE)
    parts["cmedians"].set_linewidth(2)
    ax.set_xticks([0, 1])
    ax.set_xticklabels(["Available", "Unavailable"], fontsize=8.5)
    ax.set_title(title, fontsize=10, color=NAVY, fontweight="bold", pad=4)
    ax.tick_params(labelsize=8.5)
    ax.set_facecolor(OFF_W)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, color=LGRAY, alpha=0.7)
    ax.set_axisbelow(True)

green_p = mpatches.Patch(facecolor=GREEN, alpha=0.72, label="Available next match")
red_p   = mpatches.Patch(facecolor=RED,   alpha=0.72, label="Unavailable next match")
fig2.legend(handles=[green_p, red_p], fontsize=8, loc="upper center",
            ncol=2, framealpha=0.9, bbox_to_anchor=(0.5, 1.02))
plt.tight_layout(pad=0.6, rect=[0, 0, 1, 0.95])
p2 = os.path.join(BASE, "eda_workload_distribution.png")
fig2.savefig(p2, dpi=150, bbox_inches="tight", facecolor=WHITE)
plt.close()
print(f"  Saved: {p2}")


# ── Chart 3: Absence count & rate by season ────────────────────────────────────
print("Chart 3: Absence by season...")
ss = (labeled.groupby("season", observed=True)
      .agg(total=("player_unavailable_next_match", "count"),
           absent=("player_unavailable_next_match", "sum"))
      .reset_index())
ss["available"] = ss["total"] - ss["absent"]
ss["rate"] = ss["absent"] / ss["total"]
ss["season"] = ss["season"].astype(int)

fig3, ax3 = plt.subplots(figsize=(5.5, 2.8))
fig3.patch.set_facecolor(WHITE)
ax3.set_facecolor(OFF_W)
x3 = np.arange(len(ss))
bw3 = 0.35
ax3.bar(x3 - bw3/2, ss["available"], bw3,
        color=GREEN, alpha=0.82, label="Available", edgecolor=WHITE, zorder=3)
ax3.bar(x3 + bw3/2, ss["absent"],    bw3,
        color=RED,   alpha=0.82, label="Unavailable", edgecolor=WHITE, zorder=3)
for i, (tot, ab, rate) in enumerate(zip(ss["total"], ss["absent"], ss["rate"])):
    ax3.text(i, max(tot - ab, ab) * 0.55, f"{rate*100:.0f}%",
             ha="center", va="bottom", fontsize=9.5,
             color=NAVY, fontweight="bold", zorder=5)
ax3.set_xticks(x3)
ax3.set_xticklabels([str(s) for s in ss["season"]], fontsize=10)
ax3.set_ylabel("Player-match rows", fontsize=9, color=DGRAY)
ax3.set_title("Absence Count & Rate by Season", fontsize=11, color=NAVY, fontweight="bold", pad=6)
ax3.legend(fontsize=8.5, framealpha=0.9)
ax3.grid(axis="y", zorder=0)
ax3.tick_params(labelsize=9)
ax3.spines["top"].set_visible(False)
ax3.spines["right"].set_visible(False)
plt.tight_layout(pad=0.6)
p3 = os.path.join(BASE, "eda_absence_by_season.png")
fig3.savefig(p3, dpi=150, bbox_inches="tight", facecolor=WHITE)
plt.close()
print(f"  Saved: {p3}")


# ── Chart 4: Feature correlations with target ──────────────────────────────────
print("Chart 4: Feature correlations...")
FEATURES = [
    ("minutes_last_match",          "Minutes — Last Match"),
    ("minutes_last_7_days",         "Minutes — Last 7 Days"),
    ("minutes_last_14_days",        "Minutes — Last 14 Days"),
    ("minutes_last_30_days",        "Minutes — Last 30 Days"),
    ("consecutive_starts",          "Consecutive Starts"),
    ("played_last_match",           "Played Last Match"),
    ("days_since_last_match",       "Days Since Last Match"),
    ("missed_matches_last_30_days", "Missed Matches (30d)"),
    ("days_since_last_absence",     "Days Since Last Absence"),
    ("injuries_last_30_days",       "Injuries (30d)"),
]
corrs = []
for feat, label in FEATURES:
    col = labeled[feat].fillna(labeled[feat].median())
    r = float(np.corrcoef(
        labeled["player_unavailable_next_match"].values.astype(float),
        col.values.astype(float)
    )[0, 1])
    corrs.append({"label": label, "r": r})
corr_df = pd.DataFrame(corrs).sort_values("r")

fig4, ax4 = plt.subplots(figsize=(5.5, 2.8))
fig4.patch.set_facecolor(WHITE)
ax4.set_facecolor(OFF_W)
cols4 = [RED if r > 0 else BLUE for r in corr_df["r"]]
bars4 = ax4.barh(corr_df["label"], corr_df["r"],
                  color=cols4, alpha=0.82, edgecolor=WHITE, height=0.65, zorder=3)
ax4.axvline(0, color=DGRAY, lw=1, zorder=4)
for bar, r in zip(bars4, corr_df["r"]):
    off = 0.004 if r >= 0 else -0.004
    ha  = "left" if r >= 0 else "right"
    ax4.text(r + off, bar.get_y() + bar.get_height() / 2,
             f"{r:+.3f}", va="center", ha=ha, fontsize=7.5, color=DGRAY)
ax4.set_xlabel("Point-biserial r  (binary target, continuous feature)", fontsize=8.5, color=DGRAY)
ax4.set_title("Feature Correlation with Unavailability", fontsize=11,
              color=NAVY, fontweight="bold", pad=6)
ax4.tick_params(labelsize=8)
ax4.set_xlim(-0.4, 0.4)
red_p2  = mpatches.Patch(facecolor=RED,  alpha=0.82, label="Risk ↑  (positive r)")
blue_p2 = mpatches.Patch(facecolor=BLUE, alpha=0.82, label="Risk ↓  (negative r)")
ax4.legend(handles=[red_p2, blue_p2], fontsize=8, framealpha=0.9, loc="lower right")
ax4.grid(axis="x", zorder=0)
ax4.spines["top"].set_visible(False)
ax4.spines["right"].set_visible(False)
plt.tight_layout(pad=0.6)
p4 = os.path.join(BASE, "eda_feature_correlation.png")
fig4.savefig(p4, dpi=150, bbox_inches="tight", facecolor=WHITE)
plt.close()
print(f"  Saved: {p4}")

print("\nAll 4 EDA charts saved.")
