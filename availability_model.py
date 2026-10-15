"""
Sacramento Republic FC - Player Availability Risk Model
Predicts player_unavailable_next_match (binary) using rolling workload and
absence-history features computed per player per match.
"""

import warnings
warnings.filterwarnings("ignore")

import os
import numpy as np
import pandas as pd
import openpyxl
import shap
import matplotlib.pyplot as plt
import matplotlib
matplotlib.use("Agg")

from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import (
    roc_auc_score, average_precision_score, brier_score_loss,
    classification_report, confusion_matrix, ConfusionMatrixDisplay,
    precision_recall_curve
)
from sklearn.pipeline import Pipeline
import lightgbm as lgb

BASE = r"C:\Users\fredo\500 Project"

# -- 1. Load data --------------------------------------------------------------
wb = openpyxl.load_workbook(os.path.join(BASE, "data columns v2.xlsx"))
ws = wb["Row per Player"]

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
rows = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)]

df = pd.DataFrame(rows, columns=headers)

# Fix types
df["match_date"] = pd.to_datetime(df["match_date"])
df["played_last_match"] = df["played_last_match"].map(
    {True: 1, False: 0, 1: 1, 0: 0, None: np.nan}
)
df["season"] = df["match_date"].dt.year

# -- 2. Feature set -------------------------------------------------------------
FEATURES = [
    "minutes_last_match",
    "minutes_last_7_days",
    "minutes_last_14_days",
    "minutes_last_30_days",
    "matches_last_14_days",
    "consecutive_starts",
    "played_last_match",
    "days_since_last_match",
    "missed_matches_last_30_days",
    "days_since_last_absence",
    "injuries_last_30_days",
    "injuries_last_60_days",
]
TARGET = "player_unavailable_next_match"

POSITION_DEPTH = {
    0: 1, 1: 1,
    2: 4, 4: 4, 7: 4, 8: 4, 9: 4,
    3: 0,
    5: 1, 6: 1,
    10: 1, 11: 1,
    12: 3, 13: 3, 14: 3, 15: 3,
    16: 3, 19: 3,
    17: 2, 21: 2, 23: 2,
    18: 2,
    20: 1, 22: 1,
}
df["position_depth"] = df["player_id"].map(POSITION_DEPTH)
FEATURES.append("position_depth")

df = pd.get_dummies(df, columns=["position"], prefix="pos", drop_first=False)
pos_cols = [c for c in df.columns if c.startswith("pos_")]
FEATURES += pos_cols

# -- 3. Filter to labeled rows with complete features ---------------------------
labeled = df[df[TARGET].notna()].copy()
labeled[TARGET] = labeled[TARGET].astype(int)
labeled["days_since_last_absence"] = labeled["days_since_last_absence"].fillna(999)
labeled = labeled.dropna(subset=[
    "minutes_last_match", "minutes_last_7_days",
    "minutes_last_14_days", "minutes_last_30_days",
    "days_since_last_match", "played_last_match"
])

print(f"Labeled rows after dropping NaN features: {len(labeled)}")
print(f"Target balance: {labeled[TARGET].value_counts().to_dict()}")
print(f"Seasons: {sorted(labeled['season'].unique())}")

# -- 4. Time-based train / test split ------------------------------------------
train = labeled[labeled["season"] == 2024]
test  = labeled[labeled["season"] >= 2025]

X_train = train[FEATURES].values.astype(float)
y_train = train[TARGET].values
X_test  = test[FEATURES].values.astype(float)
y_test  = test[TARGET].values

print(f"\nTrain (2024): {len(train)} rows  |  Test (2025+): {len(test)} rows")
print(f"Train positive rate: {y_train.mean()*100:.1f}%")
print(f"Test  positive rate: {y_test.mean()*100:.1f}%")

# -- 5. Logistic Regression baseline -------------------------------------------
lr_pipe = Pipeline([
    ("scaler", StandardScaler()),
    ("lr", LogisticRegression(
        class_weight="balanced",
        max_iter=1000,
        C=0.1,
        random_state=42
    ))
])
lr_pipe.fit(X_train, y_train)
lr_proba = lr_pipe.predict_proba(X_test)[:, 1]

print("\n-- Logistic Regression ----------------------------------")
print(f"ROC-AUC:   {roc_auc_score(y_test, lr_proba):.4f}")
print(f"PR-AUC:    {average_precision_score(y_test, lr_proba):.4f}")
print(f"Brier:     {brier_score_loss(y_test, lr_proba):.4f}")
print(classification_report(y_test, (lr_proba >= 0.35).astype(int),
                            target_names=["Available", "Unavailable"]))

# -- 6. LightGBM primary model -------------------------------------------------
scale_pos = (y_train == 0).sum() / (y_train == 1).sum()

lgbm = lgb.LGBMClassifier(
    n_estimators=500,
    learning_rate=0.05,
    max_depth=4,
    num_leaves=15,
    subsample=0.8,
    colsample_bytree=0.8,
    scale_pos_weight=scale_pos,
    random_state=42,
    verbose=-1
)
lgbm.fit(
    X_train, y_train,
    eval_set=[(X_test, y_test)],
    callbacks=[lgb.early_stopping(50, verbose=False), lgb.log_evaluation(0)]
)
lgb_proba = lgbm.predict_proba(X_test)[:, 1]

print("\n-- LightGBM ---------------------------------------------")
print(f"ROC-AUC:   {roc_auc_score(y_test, lgb_proba):.4f}")
print(f"PR-AUC:    {average_precision_score(y_test, lgb_proba):.4f}")
print(f"Brier:     {brier_score_loss(y_test, lgb_proba):.4f}")
print(classification_report(y_test, (lgb_proba >= 0.35).astype(int),
                            target_names=["Available", "Unavailable"]))

# -- 7. Confusion matrix plots -------------------------------------------------
fig, axes = plt.subplots(1, 2, figsize=(11, 4))
for ax, proba, title in zip(axes,
                             [lr_proba, lgb_proba],
                             ["Logistic Regression", "LightGBM"]):
    cm = confusion_matrix(y_test, (proba >= 0.35).astype(int))
    ConfusionMatrixDisplay(cm, display_labels=["Available", "Unavailable"]).plot(
        ax=ax, colorbar=False, cmap="Blues"
    )
    ax.set_title(title)
plt.tight_layout()
plt.savefig(os.path.join(BASE, "confusion_matrices.png"), dpi=150)
plt.close()
print("\nSaved: confusion_matrices.png")

# -- 8. PR curves --------------------------------------------------------------
fig, ax = plt.subplots(figsize=(7, 5))
for proba, label, color in [
    (lr_proba,  "Logistic Regression", "steelblue"),
    (lgb_proba, "LightGBM",            "darkorange"),
]:
    prec, rec, _ = precision_recall_curve(y_test, proba)
    auprc = average_precision_score(y_test, proba)
    ax.plot(rec, prec, label=f"{label}  (AUPRC={auprc:.3f})", color=color, lw=2)
baseline = y_test.mean()
ax.axhline(baseline, ls="--", color="grey", label=f"Baseline ({baseline:.2f})")
ax.set_xlabel("Recall"); ax.set_ylabel("Precision")
ax.set_title("Precision-Recall Curve")
ax.legend(); ax.grid(alpha=0.3)
plt.tight_layout()
plt.savefig(os.path.join(BASE, "pr_curves.png"), dpi=150)
plt.close()
print("Saved: pr_curves.png")

# -- 9. SHAP - global feature importance (LightGBM) ----------------------------
explainer = shap.TreeExplainer(lgbm)
shap_vals = explainer.shap_values(X_test)

if isinstance(shap_vals, list):
    sv = shap_vals[1]
else:
    sv = shap_vals

fig, ax = plt.subplots(figsize=(8, 6))
shap.summary_plot(sv, X_test, feature_names=FEATURES,
                  plot_type="bar", show=False, max_display=15)
plt.title("SHAP Feature Importance - LightGBM (Top 15)")
plt.tight_layout()
plt.savefig(os.path.join(BASE, "shap_importance.png"), dpi=150, bbox_inches="tight")
plt.close()
print("Saved: shap_importance.png")

# -- 10. SHAP beeswarm ---------------------------------------------------------
fig, ax = plt.subplots(figsize=(9, 7))
shap.summary_plot(sv, X_test, feature_names=FEATURES, show=False, max_display=15)
plt.title("SHAP Summary - Risk Drivers & Mitigators")
plt.tight_layout()
plt.savefig(os.path.join(BASE, "shap_beeswarm.png"), dpi=150, bbox_inches="tight")
plt.close()
print("Saved: shap_beeswarm.png")

# -- 11. Risk tier scoring for current squad (2026 matches) --------------------
current = df[df["season"] == 2026].copy()
current["days_since_last_absence"] = current["days_since_last_absence"].fillna(999)
current = current.dropna(subset=FEATURES).reset_index(drop=True)

if len(current) > 0:
    X_current = current[FEATURES].values.astype(float)
    current["risk_score"] = lgbm.predict_proba(X_current)[:, 1]
    current["risk_tier"] = pd.cut(
        current["risk_score"],
        bins=[-0.001, 0.15, 0.35, 1.001],
        labels=["Low (<15%)", "Medium (15-35%)", "High (>35%)"]
    )

    latest = (current.sort_values("match_date", ascending=False)
                      .groupby("player_id").first()
                      .reset_index()
                      [["player_id", "player_name", "match_date",
                        "risk_score", "risk_tier"]])
    latest = latest.sort_values("risk_score", ascending=False)

    print("\n-- Current Squad Risk Scores (most recent 2026 match) -------")
    print(latest[["player_name", "match_date", "risk_score", "risk_tier"]].to_string(index=False))

    wb2 = openpyxl.load_workbook(os.path.join(BASE, "data columns v2.xlsx"))
    if "Risk Scores" in wb2.sheetnames:
        del wb2["Risk Scores"]
    ws2 = wb2.create_sheet("Risk Scores")
    ws2.append(["player_id", "player_name", "as_of_date", "risk_score", "risk_tier"])
    for _, row in latest.iterrows():
        ws2.append([int(row["player_id"]), row["player_name"],
                    row["match_date"], round(float(row["risk_score"]), 4),
                    str(row["risk_tier"])])
    wb2.save(os.path.join(BASE, "data columns v2.xlsx"))
    print("\nRisk Scores sheet written to data columns v2.xlsx")

# -- 12. SHAP waterfall plots per player (current squad) -----------------------
if len(current) > 0:
    waterfall_dir = os.path.join(BASE, "waterfall_plots")
    os.makedirs(waterfall_dir, exist_ok=True)

    # Compute SHAP values for current squad
    sv_cur = explainer.shap_values(X_current)
    if isinstance(sv_cur, list):
        sv_cur = sv_cur[1]
    ev = explainer.expected_value
    if isinstance(ev, (list, np.ndarray)):
        ev = ev[1]

    # Map each player_id to the index of their most recent row in current
    pid_to_idx = {}
    for pid in current["player_id"].unique():
        rows_p = current[current["player_id"] == pid].sort_values("match_date", ascending=False)
        pid_to_idx[int(pid)] = rows_p.index[0]

    print(f"\n-- Waterfall plots ({len(latest)} players) -------------------")
    for _, row in latest.iterrows():
        pid  = int(row["player_id"])
        name = row["player_name"]
        risk = float(row["risk_score"])
        tier = str(row["risk_tier"])
        idx  = pid_to_idx[pid]

        exp_i = shap.Explanation(
            values=sv_cur[idx],
            base_values=float(ev),
            data=X_current[idx],
            feature_names=FEATURES
        )

        shap.plots.waterfall(exp_i, max_display=12, show=False)
        fig = plt.gcf()
        fig.suptitle(f"{name}  -  Unavailability Risk: {risk*100:.1f}%  ({tier})",
                     fontsize=11, y=1.02)
        safe = name.replace(" ", "_").replace("'", "").replace("/", "-")
        path = os.path.join(waterfall_dir, f"waterfall_{safe}.png")
        fig.savefig(path, dpi=150, bbox_inches="tight")
        plt.close(fig)
        print(f"  Saved: waterfall_{safe}.png  [{tier}]")

    # Combined grid: top 6 highest-risk players (3x2 layout)
    top6 = latest.head(6)
    fig, axes = plt.subplots(3, 2, figsize=(16, 18))
    axes_flat = axes.flatten()

    for ax_i, (_, row) in enumerate(top6.iterrows()):
        pid  = int(row["player_id"])
        name = row["player_name"]
        risk = float(row["risk_score"])
        tier = str(row["risk_tier"])
        idx  = pid_to_idx[pid]

        exp_i = shap.Explanation(
            values=sv_cur[idx],
            base_values=float(ev),
            data=X_current[idx],
            feature_names=FEATURES
        )

        plt.sca(axes_flat[ax_i])
        shap.plots.waterfall(exp_i, max_display=8, show=False)
        axes_flat[ax_i].set_title(f"{name}  -  {risk*100:.1f}% ({tier})",
                                   fontsize=9, pad=4)

    # Hide any unused panels
    for ax_i in range(len(top6), len(axes_flat)):
        axes_flat[ax_i].set_visible(False)

    fig.suptitle("Top 6 Highest-Risk Players - SHAP Waterfall", fontsize=13, y=1.01)
    plt.tight_layout()
    fig.savefig(os.path.join(BASE, "shap_waterfall_top6.png"), dpi=150, bbox_inches="tight")
    plt.close(fig)
    print(f"\nSaved: shap_waterfall_top6.png  (top 6 combined grid)")
    print(f"All waterfall plots saved to: {waterfall_dir}")

print("\nDone.")
